# -*- coding: utf-8 -*-
"""
[운영] 자동변환_소스데이터.xlsx 의 보완 사본을 만든다.

원본은 건드리지 않는다. 사본에만 반영하고, 추가된 셀은 노란 배경으로 표시한다.
판단이 갈리는 표기는 시트 '추가검토'에 따로 모아 리즈가 결정하도록 남긴다.

기준 데이터: 2026-08-05 / 07 / 10 / 11 / 12 품목표 데이터 14,463행
"""

import os
import glob
import json

import openpyxl
from openpyxl.styles import PatternFill, Font

from normalize import Master, _load_rows, clean, SHEETS, COLUMN_OF

SRC = "[운영] 자동변환_소스데이터.xlsx"
DST = "[운영] 자동변환_소스데이터_보완.xlsx"

HILITE = PatternFill("solid", fgColor="FFF3B0")
HDR = Font(bold=True)

# ---------------------------------------------------------------------------
# 1) 기존 정식명에 붙일 변형표기
#    근거: 스캔 오독이 명백하거나, 부속어만 붙은 형태
# ---------------------------------------------------------------------------

ADD_VARIANTS = {
    "warehouse": {
        "강동1": ["강한1", "강강1", "강너1"],
        "강동2": ["강강2"],
        "삼진1": ["상진1", "성진1", "삼정1"],
        "삼진2": ["상진2", "삼전2"],
        "ACE 기흥": ["기흥", "기홍", "아주기홍", "아주기훙", "에이스기퉁", "에스기흥"],
        "ACE 처인": ["처인", "치인", "아주치인", "에이스치인", "에이스치원", "에이스치언"],
        "ACE 용인": ["에이소용인", "에이소웅인", "에이쓰웅인"],
        "HL 곤지암": ["곤지암", "근지암", "곰지암", "한라근직암", "한라군지암",
                    "한라고지암", "한라근지암", "한라곶지암", "한라진지암"],
        "HL 동탄": ["동탄한라"],
        "대재": ["대체", "대제", "대계"],
        "희창": ["회창"],
        "이스트벨리": ["이스트밸리", "이스트발리"],
        "제니스곤지암": ["제니스"],
        "오로라": ["오르라"],
        "마장동": ["마장동매장"],
    },
    "item": {
        "볼라전각": ["불라전각", "블라전각", "볼갈전각", "불란전각", "불갈전각", "불라"],
        "볼살": ["불살"],
        # 리즈 확인: 부채살은 아바니코와 무관. 아바니코 행에서 떼어내 부채로 붙인다.
        "부채": ["부채살"],
        "늑간": ["늑간살"],
        "토시": ["토시살"],
        "안창": ["안창살"],
        "살치": ["살치살"],
        "홍두깨": ["홍두깨살"],
        "우둔": ["우돈", "우도"],
        "깐양": ["찐양", "꼬양", "간양", "칸양"],
        "삼겹양지": ["삼겹앞지"],
        "립엔드": ["립앤드", "RIB END"],
        "등갈비": ["뼈립", "뺑립"],
        "#2등갈비": ["백립 #2", "백립#2"],
        "목심": ["육심"],
        "목살": ["돈목살"],
        "삼겹": ["돈삼겹"],
        "전지": ["돈전지"],
        "목전지": ["돈목전지"],
        "장족": ["돈장족"],
        "단족": ["돈단족"],
    },
    "brand": {
        "STANBROKE": ["STANVROKE", "스탠브록", "스탠브룩"],
        "KILCOY": ["KICIOY", "킬코이"],
        "OAKEY": ["OAKEYS", "OKAEY", "OKAY", "OAKKYS", "오끼"],
        "SMITHFIELD": ["스미스"],
        "NOLAN": ["놀란"],
    },
    "species": {
        "우육": ["냉장우육", "냉동우육", "수입냉동우육", "육우"],
        "돈육": ["냉장돈육", "냉동돈육", "수입냉동돈육"],
        "양육": ["냉동양고기"],
        "계육": ["가금육"],
    },
}

# ---------------------------------------------------------------------------
# 2) 새로 만들 정식명 (마스터에 아예 없던 실제 품목/창고/브랜드)
# ---------------------------------------------------------------------------

NEW_ROWS = {
    "item": [
        ["목뼈"],
        ["앞목심", "압목심"],
        ["막창", "막창돈막창", "막창 돈막창"],
        ["이겹", "이겸"],
        ["갈비살"],
        ["꽃갈비"],
        ["꽃등심"],
        ["사각갈비"],
        ["앞다리살"],
        ["어깨살"],
        ["숄더랙", "숄더랙(Shoulder Rack)"],
        ["등심덧살"],
        ["앞치마살"],
        ["플랩"],
        ["리브렛"],
        ["립캡"],
        ["프렌치랙", "프렌치드랙"],
        ["앞쪽사골"],
        ["뒷사골", "뒷쪽사골"],
        ["작업항정"],
        ["스윗브레드"],
        ["찜갈비"],
        ["차돌백이"],
        ["미추리", "미추리삼겹", "미추리삼겹살"],
    ],
    "warehouse": [
        ["SM로지스", "sm로지스", "SM 로직스", "에스엠로지스"],
        # 리즈 확인: 대청냉장은 실재 창고. 대재 오독이 아니다.
        ["대청", "대치"],
    ],
    "brand": [
        ["UPPER IOWA", "어퍼아이오와", "UPPER IOWA[어퍼아이오와]"],
        ["CARGILL", "카길"],
        ["FRIGOSORNO", "프리고오소르노", "프리고소르노"],
        ["HKSCAN"],
        ["FREEREIN", "프리레인"],
        ["그로스파트너"],
        ["수라돈"],
        ["댐코타"],
        ["슈바인"],
        ["골든듀록"],
        ["잉카룹사"],
        ["그래버파크"],
        ["미트스토리"],
        ["데니시크라운", "테니시"],
        ["STERLING SILVER"],
        ["MERAMIST"],
        ["WINGHAM", "윙햄"],
        ["DELAVI", "델라비"],
        ["REIXACH", "레이삭"],
    ],
}

# ---------------------------------------------------------------------------
# 3) 마스터에 넣지 않을 것 — 창고/브랜드/품목이 아닌 값
#    (파서가 열을 잘못 넣은 경우. 이건 OCR 단계에서 잡아야 한다)
# ---------------------------------------------------------------------------

# 특정 정식명 행에서 제거할 잘못 등록된 변형표기
REMOVE_VARIANTS = {
    "item": {"아바니코": ["부채살"]},
}

NOT_A_VALUE = {
    # 상태·안내 문구가 창고 열에 들어간 것
    "예정", "문의", "통관예정", "입고예정", "입항예정", "판매완료", "판매중",
    "매주통관", "품절", "종료", "집중", "일반", "진공", "개별", "중", "x", "-",
    "NEW 집중판매", "발주후 생산", "물류예정", "인견예정", "오픈대창",
    # 등급·규격이 브랜드 열에 들어간 것
    "CH", "UN", "EX", "PR", "SE", "CAB", "S", "A", "GF", "SW", "RRA", "CS",
    "SWC", "SPLIT", "BP", "VP", "YP", "BEEF", "GOLD", "BLACK", "TF", "3P",
    "1P 진공", "CH/PR", "5RIB", "LOIN", "86K",
    # 원산지·보관이 다른 열에 들어간 것
    "호주", "미국", "스페인", "칠레", "냉장", "냉동", "와규", "듀록", "육", "도육",
    "토육", "식육", "우육", "냉장와규", "냉동돈육", "냉동우육", "냉장돈육",
    "수입냉동우육", "동결 우육", "미국산",
}


def category_of_sheet(sheet_title):
    for cat, title in SHEETS.items():
        if title == sheet_title:
            return cat
    return None


def collect_unknowns(master, min_count=2):
    rows = []
    seen = set()
    paths = sorted(glob.glob("/mnt/user-data/uploads/visionmeat/database/2026*.xlsx"))
    paths += ["20260812_품목표_데이터.xlsx"]
    for p in paths:
        b = os.path.basename(p)
        if b in seen or "회원" in b:
            continue
        seen.add(b)
        try:
            rows += _load_rows(p)
        except Exception:
            pass
    return master.unknowns(rows, min_count=min_count), len(rows)


def main():
    master = Master(SRC, cache=False)
    unk, n_rows = collect_unknowns(master)

    # 원본은 구글시트 IMPORTRANGE 수식 + 캐시값 구조라서 data_only 로 값만 읽어
    # 수식 없는 새 통합문서로 다시 쓴다. (수식 채로 두면 편집이 덮어써진다)
    src = openpyxl.load_workbook(SRC, data_only=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for s in src.worksheets:
        ws = wb.create_sheet(s.title)
        for row in s.iter_rows(values_only=True):
            vals = [clean(c) for c in row]
            while vals and not vals[-1]:
                vals.pop()
            if vals:
                ws.append(vals)
    src.close()

    added = {"variant": 0, "new": 0}
    planned = set()          # 이번에 반영한 표기 (검토목록에서 제외용)

    for ws in wb.worksheets:
        cat = category_of_sheet(ws.title)
        if not cat:
            continue

        # 정식명 -> 행번호
        rowno = {}
        for i, row in enumerate(ws.iter_rows(min_row=1, max_col=1, values_only=True), 1):
            v = clean(row[0])
            if v:
                rowno[v] = i

        # (0) 잘못 등록된 변형표기 제거
        for canon, drops in REMOVE_VARIANTS.get(cat, {}).items():
            r = rowno.get(canon)
            if not r:
                continue
            keep = [clean(c.value) for c in ws[r]
                    if clean(c.value) and clean(c.value) not in drops]
            for c in ws[r]:
                c.value = None
            for j, v in enumerate(keep, 1):
                ws.cell(row=r, column=j, value=v)
            print("  [제거] %s '%s' 에서 %s" % (ws.title, canon, ", ".join(drops)))

        # (1) 기존 행에 변형표기 추가
        for canon, variants in ADD_VARIANTS.get(cat, {}).items():
            r = rowno.get(canon)
            if not r:
                print("  [건너뜀] %s 시트에 정식명 '%s' 없음" % (ws.title, canon))
                continue
            existing = {clean(c.value) for c in ws[r] if clean(c.value)}
            col = max((c.column for c in ws[r] if clean(c.value)), default=1)
            for v in variants:
                planned.add((cat, v))
                if v in existing:
                    continue
                col += 1
                cell = ws.cell(row=r, column=col, value=v)
                cell.fill = HILITE
                added["variant"] += 1

        # (2) 새 정식명 행 추가
        last = ws.max_row
        for vals in NEW_ROWS.get(cat, []):
            if clean(vals[0]) in rowno:
                continue
            last += 1
            for j, v in enumerate(vals, 1):
                planned.add((cat, v))
                cell = ws.cell(row=last, column=j, value=v)
                cell.fill = HILITE
            added["new"] += 1

    # (3) 판단 보류분 -> 검토 시트
    if "추가검토" in wb.sheetnames:
        del wb["추가검토"]
    rv = wb.create_sheet("추가검토")
    rv.append(["구분", "발견된 표기", "횟수", "자동추정", "메모"])
    for c in rv[1]:
        c.font = HDR
        c.fill = HILITE

    n_review = 0
    for cat in ("item", "brand", "warehouse", "species"):
        for v, n, guess in unk.get(cat, []):
            if (cat, v) in planned:
                continue
            memo = ""
            if v in NOT_A_VALUE:
                memo = "창고/브랜드/품목 아님 - 열 오배치 의심"
            rv.append([SHEETS[cat], v, n, guess, memo])
            n_review += 1

    for col, w in zip("ABCDE", (18, 34, 8, 16, 34)):
        rv.column_dimensions[col].width = w

    wb.save(DST)

    print("\n기준 데이터: %d행" % n_rows)
    print("변형표기 추가: %d건" % added["variant"])
    print("신규 정식명 추가: %d건" % added["new"])
    print("추가검토 시트: %d건" % n_review)
    print("저장: %s" % DST)


if __name__ == "__main__":
    main()
